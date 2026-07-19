"""Exportación de reportes a CSV, Excel y PDF (RF-WEB-11.2)."""
import csv
import io


def _columns(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in cols:
                cols.append(k)
    return cols


def to_csv(title: str, rows: list[dict]) -> bytes:
    buf = io.StringIO()
    cols = _columns(rows)
    writer = csv.DictWriter(buf, fieldnames=cols)
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")  # BOM para Excel


def to_xlsx(title: str, rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Reporte"
    cols = _columns(rows)
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = header_fill
        cell.font = header_font
    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(cols, start=1):
            ws.cell(row=r, column=c, value=row.get(name))
    for c, name in enumerate(cols, start=1):
        width = max([len(str(name))] + [len(str(row.get(name, ""))) for row in rows]) + 2
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = min(width, 40)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(title: str, rows: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4),
                            leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 0.5 * cm)]

    cols = _columns(rows)
    if not rows:
        elements.append(Paragraph("Sin datos para los filtros seleccionados.", styles["Normal"]))
    else:
        data = [cols] + [[str(row.get(c, "")) for c in cols] for row in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(table)
    doc.build(elements)
    return out.getvalue()


EXPORTERS = {
    "csv": (to_csv, "text/csv", "csv"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "pdf": (to_pdf, "application/pdf", "pdf"),
}
